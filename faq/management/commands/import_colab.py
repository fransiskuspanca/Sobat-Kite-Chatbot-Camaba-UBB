"""Artefak notebook -> basis data.

Menerima dua bentuk masukan:
  1. knowledge_base.json  — keluaran `df.to_json(orient="records")` di notebook
  2. Excel panitia PMB    — diparse dengan parser YANG SAMA seperti notebook

Parser Excel dan aturan flag `butuh_verifikasi` disalin dari notebook (sel 2
dan 3b) supaya isi basis data identik dengan yang dipakai saat melatih model.

Variasi cara bertanya ikut diimpor bila ada: dari sheet "Variasi" di Excel
(kolom "Pertanyaan utama" + "Variasi"), atau dari kunci "variasi" di JSON
keluaran backup_kb. Variasi yang sudah ada di basis data tidak dihapus.
"""

import json
import re
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from faq.models import Faq, FaqVariasi

# --- disalin dari notebook, sel 2 ----------------------------------------
KATEGORI_RE = re.compile(r"^[A-Z]{1,2}\.\s+\S")
HEADER_TOKENS = {"no", "pertanyaan", "jawaban", "sumber"}

# --- disalin dari notebook, sel 3b ---------------------------------------
PLACEHOLDER_PAT = r"perlu wawancara|belum ditemukan informasi|tidak ditemukan informasi"


def parse_faq_excel(path, sheet=0):
    """Parser stateful untuk Excel FAQ berstruktur seksi.

    Mengembalikan (list_of_dict, list_of_warnings).
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[sheet]] if isinstance(sheet, int) else wb[sheet]

    rows, warns = [], []
    kategori, sub_kategori = None, None

    for i, raw in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = [("" if c is None else str(c).strip()) for c in raw]
        cells += [""] * (4 - len(cells))          # pastikan minimal 4 kolom
        no, pertanyaan, jawaban, sumber = cells[:4]

        if not any(cells):                        # baris pemisah kosong
            continue

        # Baris judul: hanya kolom A yang terisi
        if no and not pertanyaan and not jawaban:
            if KATEGORI_RE.match(no):
                kategori, sub_kategori = no, None
            else:
                sub_kategori = no
            continue

        # Baris header tabel yang berulang di tiap seksi
        if no.lower() in HEADER_TOKENS and pertanyaan.lower() in HEADER_TOKENS:
            continue

        if pertanyaan and jawaban:
            rows.append({
                "kategori":     kategori or "(tanpa kategori)",
                "sub_kategori": sub_kategori or "",
                "pertanyaan":   " ".join(pertanyaan.split()),
                "jawaban":      " ".join(jawaban.split()),
                "sumber":       " ".join(sumber.split()),
                "baris_excel":  i,
            })
        elif pertanyaan and not jawaban:
            warns.append(f"Baris {i}: pertanyaan tanpa jawaban -> {pertanyaan[:70]!r}")
        elif jawaban and not pertanyaan:
            warns.append(f"Baris {i}: jawaban tanpa pertanyaan -> {jawaban[:70]!r}")

    return rows, warns


SHEET_VARIASI = "Variasi"
ASAL_DIPERTAHANKAN = ("log", "augmentasi")


def _norm(teks) -> str:
    return " ".join(str(teks or "").lower().split())


def parse_variasi_excel(path, sheet=SHEET_VARIASI):
    """Baca sheet variasi: satu baris = satu variasi untuk satu pertanyaan.

    Kolom dicari lewat judul di baris pertama ("Pertanyaan utama" dan
    "Variasi"), jadi urutan dan kolom tambahan seperti No/Kategori bebas.
    Mengembalikan (list_of_(pertanyaan, variasi), list_of_warnings); list
    kosong bila sheet tidak ada.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        return [], []
    ws = wb[sheet]

    baris = ws.iter_rows(values_only=True)
    judul = [_norm(c) for c in next(baris, ())]
    if "pertanyaan utama" not in judul or "variasi" not in judul:
        return [], [f"Sheet {sheet!r} dilewati: butuh kolom 'Pertanyaan utama' "
                    "dan 'Variasi' di baris pertama"]
    i_tanya, i_var = judul.index("pertanyaan utama"), judul.index("variasi")

    hasil, warns = [], []
    for i, raw in enumerate(baris, start=2):
        raw = list(raw) + [None] * (max(i_tanya, i_var) + 1 - len(raw))
        tanya = " ".join(str(raw[i_tanya] or "").split())
        var = " ".join(str(raw[i_var] or "").split())
        if not tanya and not var:
            continue
        if not tanya or not var:
            warns.append(f"Variasi baris {i}: kolom kosong -> {(tanya or var)[:70]!r}")
            continue
        hasil.append((tanya, var))
    return hasil, warns


def flag_butuh_verifikasi(jawaban: str) -> bool:
    """Aturan notebook sel 3b — jawaban placeholder tidak boleh tampil."""
    return bool(re.search(PLACEHOLDER_PAT, (jawaban or "").lower()))


def _pisah_sumber(sumber: str) -> tuple[str, str]:
    """Kolom 'sumber' Excel bisa berisi URL atau label bebas."""
    s = (sumber or "").strip()
    if s in {"", "—", "-"}:
        return "", ""
    m = re.search(r"https?://\S+", s)
    if m:
        url = m.group(0).rstrip(").,;")
        label = s.replace(m.group(0), "").strip(" -–—:()")
        return label[:200], url
    return s[:200], ""


class Command(BaseCommand):
    help = "Impor knowledge_base.json atau Excel panitia PMB ke tabel Faq."

    def add_arguments(self, parser):
        parser.add_argument("sumber", help="Path ke knowledge_base.json atau .xlsx")
        parser.add_argument("--sheet", default=0,
                            help="Nama atau indeks sheet (khusus Excel)")
        parser.add_argument("--hapus-dulu", action="store_true",
                            help="Kosongkan tabel Faq sebelum impor")
        parser.add_argument("--dry-run", action="store_true",
                            help="Tampilkan ringkasan tanpa menulis ke basis data")

    def handle(self, *args, **opts):
        path = Path(opts["sumber"])
        if not path.exists():
            raise CommandError(f"Berkas tidak ditemukan: {path}")

        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            sheet = opts["sheet"]
            try:
                sheet = int(sheet)
            except (TypeError, ValueError):
                pass
            baris, warns = parse_faq_excel(path, sheet)
            variasi, warns_var = parse_variasi_excel(path)
            warns += warns_var
            for w in warns[:10]:
                self.stderr.write(self.style.WARNING(f"  ⚠️  {w}"))
            if len(warns) > 10:
                self.stderr.write(self.style.WARNING(
                    f"  ... dan {len(warns) - 10} peringatan lain"))
        elif path.suffix.lower() == ".json":
            try:
                with open(path, encoding="utf-8") as f:
                    baris = json.load(f)
            except json.JSONDecodeError as e:
                raise CommandError(f"{path} bukan JSON yang sah: {e}")

            # Bentuk yang diharapkan: list of dict, keluaran backup_kb atau
            # df.to_json(orient="records") di notebook.
            if isinstance(baris, dict):
                raise CommandError(
                    f"{path} berisi objek JSON, bukan daftar. Berkas yang "
                    "benar diawali '[' dan berisi daftar FAQ — periksa apakah "
                    "isinya terbungkus kunci lain.")
            if not isinstance(baris, list):
                raise CommandError(
                    f"{path} berisi {type(baris).__name__}, seharusnya daftar FAQ.")
            variasi = [(r.get("pertanyaan", ""), v)
                       for r in baris if isinstance(r, dict)
                       for v in (r.get("variasi") or [])]
        else:
            raise CommandError(f"Format tidak dikenali: {path.suffix}")

        if not baris:
            ukuran = path.stat().st_size
            raise CommandError(
                f"{path} tidak berisi satu pun FAQ (ukuran {ukuran} byte).\n"
                "Kemungkinan berkas ini hasil ekspor dari basis data yang masih "
                "kosong. Ekspor ulang dari mesin yang datanya lengkap:\n"
                "  python manage.py backup_kb --keluar kb-awal.json\n"
                "atau impor langsung dari Excel panitia PMB.")

        # Deduplikasi berdasarkan teks pertanyaan (Bagian 7.2: entri kembar
        # membuat retrieval ambigu selamanya).
        terlihat: dict[str, dict] = {}
        kembar = []
        for r in baris:
            kunci = " ".join(str(r.get("pertanyaan", "")).lower().split())
            if not kunci:
                continue
            if kunci in terlihat:
                kembar.append(r["pertanyaan"])
                continue
            terlihat[kunci] = r

        n_flag = sum(1 for r in terlihat.values()
                     if flag_butuh_verifikasi(r.get("jawaban", "")))

        self.stdout.write(f"Baris terparsing   : {len(baris)}")
        self.stdout.write(f"Unik               : {len(terlihat)}")
        self.stdout.write(f"Kembar dilewati    : {len(kembar)}")
        self.stdout.write(f"Butuh verifikasi   : {n_flag} "
                          f"({n_flag / len(terlihat):.1%})")
        self.stdout.write(f"Variasi di berkas  : {len(variasi)}")
        for k in kembar[:5]:
            self.stdout.write(self.style.WARNING(f"  ⚠️  kembar: {k[:70]}"))

        if opts["dry_run"]:
            self.stdout.write(self.style.WARNING("\n--dry-run: tidak ada yang ditulis."))
            return

        with transaction.atomic():
            simpan = []
            if opts["hapus_dulu"]:
                # Variasi dari log camaba (tombol Tautkan) dan paket sinkron
                # hanya ada di basis data, tidak di Excel. Ikut terhapus
                # bersama FAQ-nya, jadi dipasang kembali setelah impor.
                simpan = list(FaqVariasi.objects.filter(asal__in=ASAL_DIPERTAHANKAN)
                              .values_list("faq__pertanyaan", "teks", "asal"))
                n = Faq.objects.count()
                Faq.objects.all().delete()
                self.stdout.write(self.style.WARNING(f"{n} entri lama dihapus."))

            dibuat = diperbarui = 0
            for r in terlihat.values():
                label, url = _pisah_sumber(r.get("sumber", ""))
                nilai = {
                    "jawaban": r.get("jawaban", ""),
                    "kategori": r.get("kategori", "") or "",
                    "sub_kategori": r.get("sub_kategori", "") or "",
                    "sumber_label": label,
                    "sumber_url": url,
                    "butuh_verifikasi": flag_butuh_verifikasi(r.get("jawaban", "")),
                }
                obj, baru = Faq.objects.update_or_create(
                    pertanyaan=r["pertanyaan"], defaults=nilai)
                dibuat += baru
                diperbarui += not baru

            var_baru, var_yatim = self._impor_variasi(variasi)
            dipulihkan = self._pulihkan_variasi(simpan)

        self.stdout.write(self.style.SUCCESS(
            f"\n✅ {dibuat} dibuat, {diperbarui} diperbarui. "
            f"Total FAQ: {Faq.objects.count()}, "
            f"variasi: {FaqVariasi.objects.count()} ({var_baru} baru)"))
        if simpan:
            self.stdout.write(
                f"Variasi log/sinkron dipulihkan: {dipulihkan} dari {len(simpan)} "
                "(sisanya milik FAQ yang tidak ada lagi di berkas)")
        for t in var_yatim[:5]:
            self.stderr.write(self.style.WARNING(
                f"  ⚠️  variasi tanpa FAQ: {t[:70]}"))
        if len(var_yatim) > 5:
            self.stderr.write(self.style.WARNING(
                f"  ... dan {len(var_yatim) - 5} pertanyaan utama lain tidak ditemukan"))
        self.stdout.write(
            "Catatan: entri butuh_verifikasi TETAP masuk index; jawabannya "
            "diganti pesan eskalasi saat runtime (prinsip #4).")

    def _pulihkan_variasi(self, simpan):
        faq_per_teks = {_norm(f.pertanyaan): f for f in Faq.objects.all()}
        n = 0
        for tanya, teks, asal in simpan:
            faq = faq_per_teks.get(_norm(tanya))
            if faq is None:
                continue
            _, dibuat = FaqVariasi.objects.get_or_create(
                faq=faq, teks=teks, defaults={"asal": asal})
            n += dibuat
        return n

    def _impor_variasi(self, variasi):
        """Tambahkan variasi ke FAQ yang pertanyaannya cocok (abaikan besar
        kecil huruf & spasi). Mengembalikan (jumlah_baru, pertanyaan_yatim)."""
        faq_per_teks = {_norm(f.pertanyaan): f for f in Faq.objects.all()}
        baru, yatim = 0, []
        for tanya, teks in variasi:
            faq = faq_per_teks.get(_norm(tanya))
            teks = " ".join(str(teks or "").split())
            if faq is None:
                if tanya not in yatim:
                    yatim.append(tanya)
                continue
            # Variasi yang sama dengan pertanyaannya sendiri hanya
            # menggandakan vektor yang sudah ada.
            if not teks or _norm(teks) == _norm(faq.pertanyaan):
                continue
            _, dibuat = FaqVariasi.objects.get_or_create(
                faq=faq, teks=teks, defaults={"asal": "admin"})
            baru += dibuat
        return baru, yatim
