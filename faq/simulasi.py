"""Simulasi pertanyaan tanpa camaba sungguhan.

Dua jenis uji, keduanya memakai FAQ & variasi yang ada di basis data saat ini
(bukan index aktif), sehingga hasilnya mencerminkan index yang AKAN dipasang
pada reindex berikutnya:

1. Uji variasi (leave-one-out): setiap pertanyaan kanonik dan variasi dicari
   seolah-olah ditanyakan camaba, dengan vektornya sendiri disembunyikan.
   Kalau yang menang FAQ lain, variasi itu rawan "direbut" — penyebab
   jawaban tidak nyambung yang paling sering setelah variasi ditambah massal.
2. Uji pertanyaan: daftar pertanyaan bebas (berkas, hard eval, log camaba)
   dijawab lewat empat tier seperti di API.
"""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np

from faq.models import Faq


def metadata_faq(ids: list[int]) -> dict[int, dict]:
    keluar = {}
    for f in Faq.objects.filter(id__in=ids, aktif=True):
        keluar[f.id] = dict(
            pertanyaan=f.pertanyaan, jawaban=f.jawaban, kategori=f.kategori,
            sumber_label=f.sumber_label, sumber_url=f.sumber_url,
            butuh_verifikasi=f.butuh_verifikasi,
            terakhir_diverifikasi=(f.terakhir_diverifikasi.isoformat()
                                   if f.terakhir_diverifikasi else None))
    return keluar


def siapkan_penguji(engine, entri: list[dict]):
    """Engine baru di atas snapshot dari `entri`, tanpa menyentuh index aktif."""
    from retrieval.engine import Engine
    from retrieval.indexer import bangun_snapshot

    from faq.services import buat_penjaga

    snap = bangun_snapshot(engine, entri, versi=0)
    penguji = Engine(engine.model, snap, tau_high=engine.tau_high,
                     tau_low=engine.tau_low, top_k=engine.top_k,
                     query_prefix=engine.query_prefix,
                     doc_prefix=engine.doc_prefix)
    penguji.penjaga = buat_penjaga(entri)
    return penguji


def uji_variasi(penguji, entri: list[dict]) -> tuple[float, list[dict]]:
    """Leave-one-out atas seluruh entri. Kembalikan (akurasi, daftar rawan)."""
    snap = penguji.snapshot
    pemilik = np.asarray(snap.pemilik)
    q = penguji.encode([e["teks"] for e in entri], is_query=True)
    skor = q @ snap.vektor.T
    np.fill_diagonal(skor, -np.inf)          # sembunyikan vektornya sendiri

    nama = dict(Faq.objects.values_list("id", "pertanyaan"))
    rawan = []
    for i, e in enumerate(entri):
        baris = skor[i]
        milik = baris[pemilik == e["faq_id"]]
        skor_milik = float(milik.max()) if milik.size and np.isfinite(milik.max()) else None
        j = int(np.argmax(baris))
        if pemilik[j] == e["faq_id"]:
            continue
        rawan.append({
            "teks": e["teks"],
            "faq_id": e["faq_id"],
            "milik": nama.get(e["faq_id"], ""),
            "dapat_faq_id": int(pemilik[j]),
            "dapat": nama.get(int(pemilik[j]), ""),
            "skor_dapat": round(float(baris[j]), 4),
            "skor_milik": round(skor_milik, 4) if skor_milik is not None else None,
        })
    rawan.sort(key=lambda r: r["skor_dapat"] - (r["skor_milik"] or 0), reverse=True)
    return 1 - len(rawan) / len(entri), rawan


def uji_pertanyaan(penguji, daftar: list[tuple[str, str | None]]) -> list[dict]:
    """Jawab setiap pertanyaan lewat empat tier. `harapan` = pertanyaan kanonik."""
    peta = {" ".join(p.lower().split()): i
            for i, p in Faq.objects.values_list("id", "pertanyaan")}
    hasil = []
    for q, harapan in daftar:
        k = penguji.jawab(q, metadata_faq)
        atas = k.kandidat[0] if k.kandidat else {}
        emas = peta.get(" ".join(harapan.lower().split())) if harapan else None
        hasil.append({
            "pertanyaan": q,
            "tipe": str(k.tipe),
            "skor": round(k.skor, 4),
            "terpilih": k.matched or atas.get("pertanyaan", ""),
            "harapan": harapan or "",
            "benar": (None if emas is None
                      else (k.faq_id or atas.get("id")) == emas),
        })
    return hasil


def baca_berkas(path: str | Path) -> list[tuple[str, str | None]]:
    """.txt (satu pertanyaan per baris), .json ({kanonik: [parafrase]} atau
    daftar), atau .xlsx/.csv (kolom pertama pertanyaan, kolom kedua harapan)."""
    p = Path(path)
    akhiran = p.suffix.lower()
    if akhiran == ".json":
        data = json.loads(p.read_text(encoding="utf-8"))
        if isinstance(data, dict):
            return [(q, k) for k, qs in data.items() for q in qs]
        return [(str(q), None) for q in data]
    if akhiran in {".xlsx", ".xlsm"}:
        import openpyxl
        ws = openpyxl.load_workbook(p, read_only=True).active
        keluar = []
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if not r or not r[0]:
                continue
            if i == 0 and str(r[0]).strip().lower() in {"pertanyaan", "query"}:
                continue
            harapan = r[1] if len(r) > 1 and r[1] else None
            keluar.append((str(r[0]).strip(), str(harapan).strip() if harapan else None))
        return keluar
    if akhiran == ".csv":
        import csv
        with open(p, encoding="utf-8", newline="") as f:
            baris = [r for r in csv.reader(f) if r and r[0].strip()]
        if baris and baris[0][0].strip().lower() in {"pertanyaan", "query"}:
            baris = baris[1:]
        return [(r[0].strip(), r[1].strip() if len(r) > 1 and r[1].strip() else None)
                for r in baris]
    return [(t.strip(), None) for t in p.read_text(encoding="utf-8").splitlines()
            if t.strip()]


def tulis_laporan(path: str | Path, ringkasan: list[tuple[str, object]],
                  rawan: list[dict], pertanyaan: list[dict]) -> None:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    for baris in ringkasan:
        ws.append(list(baris))
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18

    def lembar(judul, kolom, data):
        s = wb.create_sheet(judul)
        s.append([k for k, _ in kolom])
        for c in s[1]:
            c.font = Font(bold=True)
        for d in data:
            s.append([d.get(k) for _, k in kolom])
        for i, (_, k) in enumerate(kolom):
            s.column_dimensions[chr(65 + i)].width = 60 if k in {
                "teks", "milik", "dapat", "pertanyaan", "terpilih", "harapan"} else 12
        s.freeze_panes = "A2"
        s.auto_filter.ref = s.dimensions

    lembar("Variasi rawan", [("Variasi", "teks"), ("FAQ pemilik", "milik"),
                             ("Direbut oleh", "dapat"), ("Skor perebut", "skor_dapat"),
                             ("Skor pemilik", "skor_milik")], rawan)
    lembar("Uji pertanyaan", [("Pertanyaan", "pertanyaan"), ("Tipe", "tipe"),
                              ("Skor", "skor"), ("FAQ terpilih", "terpilih"),
                              ("Harapan", "harapan"), ("Benar", "benar")], pertanyaan)
    wb.save(path)
