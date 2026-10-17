"""Tes simulasi pertanyaan & tautan antrean konten -> variasi."""

import openpyxl
import pytest
from django.contrib.auth.models import User
from django.core.management import call_command
from django.test import Client
from django.urls import reverse

from faq import simulasi as sim
from faq.actions import kumpulkan_entri
from faq.models import Faq, FaqVariasi
from percakapan.models import ChatLog, Feedback
from retrieval.engine import Engine
from tests.test_engine import ModelTiruan

pytestmark = pytest.mark.django_db


@pytest.fixture
def kb():
    biaya = Faq.objects.create(pertanyaan="berapa biaya kuliah", jawaban="Rp1 juta.")
    daftar = Faq.objects.create(pertanyaan="kapan pendaftaran dibuka",
                                jawaban="Maret.")
    FaqVariasi.objects.create(faq=biaya, teks="biaya kuliah per semester berapa")
    # Variasi yang salah tempat: maknanya milik FAQ pendaftaran.
    FaqVariasi.objects.create(faq=biaya, teks="kapan jadwal pendaftaran dibuka")
    FaqVariasi.objects.create(faq=daftar, teks="jadwal pendaftaran dibuka kapan")
    return {"biaya": biaya, "daftar": daftar}


@pytest.fixture
def engine_tiruan(monkeypatch):
    eng = Engine(ModelTiruan(), None, tau_high=0.72, tau_low=0.60)
    import faq.services as svc
    monkeypatch.setattr(svc, "muat_engine", lambda: eng)
    return eng


def test_uji_variasi_menemukan_variasi_yang_direbut(kb, engine_tiruan):
    entri, _ = kumpulkan_entri()
    akurasi, rawan = sim.uji_variasi(sim.siapkan_penguji(engine_tiruan, entri), entri)
    teks = {r["teks"] for r in rawan}
    assert "kapan jadwal pendaftaran dibuka" in teks
    assert "biaya kuliah per semester berapa" not in teks
    r = next(r for r in rawan if r["teks"] == "kapan jadwal pendaftaran dibuka")
    assert r["dapat"] == "kapan pendaftaran dibuka"
    assert akurasi < 1


def test_uji_pertanyaan_menilai_harapan(kb, engine_tiruan):
    entri, _ = kumpulkan_entri()
    hasil = sim.uji_pertanyaan(sim.siapkan_penguji(engine_tiruan, entri),
                               [("berapa biaya kuliah", "berapa biaya kuliah"),
                                ("resep rendang padang", None)])
    assert hasil[0]["benar"] is True
    assert hasil[0]["tipe"] == "jawaban"
    assert hasil[1]["benar"] is None
    assert hasil[1]["tipe"] == "fallback"


@pytest.mark.parametrize("akhiran,isi", [
    (".txt", "kapan daftar\n\nberapa biaya\n"),
    (".json", '{"kapan pendaftaran dibuka": ["kapan daftar", "berapa biaya"]}'),
    (".csv", "pertanyaan,harapan\nkapan daftar,kapan pendaftaran dibuka\nberapa biaya,\n"),
])
def test_baca_berkas(tmp_path, akhiran, isi):
    p = tmp_path / f"uji{akhiran}"
    p.write_text(isi, encoding="utf-8")
    assert [q for q, _ in sim.baca_berkas(p)] == ["kapan daftar", "berapa biaya"]


def test_baca_berkas_xlsx(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["Pertanyaan", "Harapan"])
    wb.active.append(["kapan daftar", "kapan pendaftaran dibuka"])
    wb.save(tmp_path / "uji.xlsx")
    assert sim.baca_berkas(tmp_path / "uji.xlsx") == [
        ("kapan daftar", "kapan pendaftaran dibuka")]


def test_perintah_simulasi_menulis_laporan(kb, engine_tiruan, tmp_path):
    berkas = tmp_path / "uji.txt"
    berkas.write_text("berapa biaya kuliah\nresep rendang padang\n", encoding="utf-8")
    ChatLog.objects.create(log_id="L1", query="kapan wisuda", tipe="fallback", skor=0.3)
    keluar = tmp_path / "laporan.xlsx"
    call_command("simulasi", "--berkas", str(berkas), "--log-hari", "7",
                 "--keluar", str(keluar))
    wb = openpyxl.load_workbook(keluar)
    assert wb.sheetnames == ["Ringkasan", "Variasi rawan", "Uji pertanyaan"]
    uji = [r[0] for r in wb["Uji pertanyaan"].iter_rows(min_row=2, values_only=True)]
    assert {"berapa biaya kuliah", "resep rendang padang", "kapan wisuda"} <= set(uji)


# --- antrean konten -> tautkan ke FAQ ---------------------------------------

@pytest.fixture
def klien_admin():
    c = Client()
    c.force_login(User.objects.create_superuser("op", "op@ubb.ac.id", "rahasia"))
    return c


@pytest.fixture
def engine_siap_tiruan(kb, monkeypatch):
    entri, _ = kumpulkan_entri()
    eng = sim.siapkan_penguji(Engine(ModelTiruan(), None), entri)
    import faq.services as svc
    monkeypatch.setattr(svc, "engine_siap", lambda: eng)
    return eng


def test_antrean_menampilkan_saran_faq(klien_admin, engine_siap_tiruan):
    ChatLog.objects.create(log_id="L1", query="biaya kuliah berapa ya",
                           tipe="fallback", skor=0.5)
    isi = klien_admin.get(reverse("admin:antrean_konten")).content.decode()
    assert "Tautkan" in isi
    assert "berapa biaya kuliah" in isi


def test_antrean_tanpa_engine_tetap_terbuka(klien_admin, kb, monkeypatch):
    import faq.services as svc
    monkeypatch.setattr(svc, "engine_siap", lambda: None)
    ChatLog.objects.create(log_id="L1", query="biaya kuliah berapa ya",
                           tipe="fallback", skor=0.5)
    r = klien_admin.get(reverse("admin:antrean_konten"))
    assert r.status_code == 200
    assert "biaya kuliah berapa ya" in r.content.decode()


def test_jawaban_jempol_bawah_masuk_antrean(klien_admin, kb, monkeypatch):
    import faq.services as svc
    monkeypatch.setattr(svc, "engine_siap", lambda: None)
    log = ChatLog.objects.create(log_id="L1", query="ukt mahal gak",
                                 tipe="jawaban", skor=0.8)
    Feedback.objects.create(log=log, rating="down")
    ChatLog.objects.create(log_id="L2", query="kapan daftar", tipe="jawaban", skor=0.9)
    isi = klien_admin.get(reverse("admin:antrean_konten")).content.decode()
    assert "ukt mahal gak" in isi
    assert "kapan daftar" not in isi


def test_tautkan_membuat_variasi_dan_hilang_dari_antrean(klien_admin, kb, monkeypatch):
    import faq.services as svc
    monkeypatch.setattr(svc, "engine_siap", lambda: None)
    ChatLog.objects.create(log_id="L1", query="Biaya  kuliah berapa ya",
                           tipe="fallback", skor=0.5)
    r = klien_admin.post(reverse("admin:antrean_tautkan"),
                         {"query": "Biaya  kuliah berapa ya", "faq_id": kb["biaya"].pk})
    assert r.status_code == 302
    v = FaqVariasi.objects.get(teks="Biaya kuliah berapa ya")
    assert v.faq == kb["biaya"] and v.asal == "log"
    isi = klien_admin.get(reverse("admin:antrean_konten")).content.decode()
    assert "Biaya  kuliah berapa ya" not in isi
    # Klik dua kali tidak membuat variasi kembar.
    klien_admin.post(reverse("admin:antrean_tautkan"),
                     {"query": "Biaya kuliah berapa ya", "faq_id": kb["biaya"].pk})
    assert FaqVariasi.objects.filter(teks="Biaya kuliah berapa ya").count() == 1


def test_tautkan_menolak_tanpa_izin(kb):
    u = User.objects.create_user("tamu", password="x", is_staff=True)
    c = Client()
    c.force_login(u)
    c.post(reverse("admin:antrean_tautkan"),
           {"query": "biaya kuliah", "faq_id": kb["biaya"].pk})
    assert not FaqVariasi.objects.filter(teks="biaya kuliah").exists()


def test_tautkan_faq_tidak_ada(klien_admin, kb):
    r = klien_admin.post(reverse("admin:antrean_tautkan"),
                         {"query": "biaya kuliah", "faq_id": 99999})
    assert r.status_code == 302
    assert not FaqVariasi.objects.filter(teks="biaya kuliah").exists()
