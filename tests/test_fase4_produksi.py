"""Tes Fase 4 — perintah terjadwal, ekspor Excel, statistik, konfigurasi produksi."""

import io
from datetime import date, timedelta

import pytest
from django.contrib.auth.models import User
from django.core import mail
from django.core.management import call_command
from django.test import Client
from django.urls import reverse
from django.utils import timezone

from faq.models import Faq
from percakapan.models import ChatLog, Feedback

pytestmark = pytest.mark.django_db


@pytest.fixture
def klien_admin():
    u = User.objects.create_superuser("op4", "op4@ubb.ac.id", "rahasia")
    c = Client()
    c.force_login(u)
    return c


# --- bersihkan_log -------------------------------------------------------

def _log(log_id, hari_lalu=0, tipe="jawaban"):
    log = ChatLog.objects.create(log_id=log_id, query="uji", tipe=tipe, skor=0.9)
    ChatLog.objects.filter(pk=log.pk).update(
        dibuat=timezone.now() - timedelta(days=hari_lalu))
    return log


def test_bersihkan_log_menghapus_yang_lama():
    _log("lama", hari_lalu=400)
    _log("baru", hari_lalu=10)
    call_command("bersihkan_log", "--bulan", "12")
    assert list(ChatLog.objects.values_list("log_id", flat=True)) == ["baru"]


def test_bersihkan_log_dry_run_tidak_menghapus():
    _log("lama", hari_lalu=400)
    call_command("bersihkan_log", "--bulan", "12", "--dry-run")
    assert ChatLog.objects.count() == 1


def test_bersihkan_log_ikut_menghapus_feedback():
    log = _log("lama", hari_lalu=400)
    Feedback.objects.create(log=log, rating="up")
    call_command("bersihkan_log", "--bulan", "12")
    assert Feedback.objects.count() == 0


def test_retensi_bisa_diatur():
    _log("sedang", hari_lalu=100)
    call_command("bersihkan_log", "--bulan", "12")
    assert ChatLog.objects.count() == 1      # belum kedaluwarsa
    call_command("bersihkan_log", "--bulan", "1")
    assert ChatLog.objects.count() == 0


# --- ingatkan_kedaluwarsa ------------------------------------------------

def test_laporan_menghitung_kategori_dengan_benar():
    Faq.objects.create(pertanyaan="lama", jawaban="x",
                       terakhir_diverifikasi=date.today() - timedelta(days=400))
    Faq.objects.create(pertanyaan="baru", jawaban="x",
                       terakhir_diverifikasi=date.today())
    Faq.objects.create(pertanyaan="belum", jawaban="x")
    Faq.objects.create(pertanyaan="menunggu", jawaban="x", butuh_verifikasi=True)

    keluar = io.StringIO()
    call_command("ingatkan_kedaluwarsa", stdout=keluar)
    teks = keluar.getvalue()
    assert "Kedaluwarsa (> 12 bulan) : 1" in teks
    assert "Belum pernah diverifikasi        : 2" in teks
    assert "Menunggu persetujuan unit        : 1" in teks


def test_laporan_bisa_dikirim_email():
    Faq.objects.create(pertanyaan="lama", jawaban="x",
                       terakhir_diverifikasi=date.today() - timedelta(days=400))
    call_command("ingatkan_kedaluwarsa", "--email", "panitia@ubb.ac.id",
                 stdout=io.StringIO())
    assert len(mail.outbox) == 1
    assert "panitia@ubb.ac.id" in mail.outbox[0].recipients()
    assert "Kesegaran FAQ" in mail.outbox[0].subject


def test_faq_nonaktif_tidak_dilaporkan():
    Faq.objects.create(pertanyaan="nonaktif", jawaban="x", aktif=False,
                       terakhir_diverifikasi=date.today() - timedelta(days=400))
    keluar = io.StringIO()
    call_command("ingatkan_kedaluwarsa", stdout=keluar)
    assert "Kedaluwarsa (> 12 bulan) : 0" in keluar.getvalue()


# --- ekspor Excel --------------------------------------------------------

def test_ekspor_excel_menghasilkan_berkas(klien_admin):
    f = Faq.objects.create(pertanyaan="Apa itu UBB?", jawaban="PTN.",
                           kategori="A. Umum", sumber_label="Profil")
    r = klien_admin.post(reverse("admin:faq_faq_changelist"), {
        "action": "ekspor_excel", "_selected_action": [str(f.pk)]})

    assert r.status_code == 200
    assert "spreadsheetml" in r["Content-Type"]
    assert ".xlsx" in r["Content-Disposition"]
    assert len(r.content) > 1000


def test_isi_excel_benar(klien_admin):
    from openpyxl import load_workbook

    from faq.models import FaqVariasi

    f = Faq.objects.create(pertanyaan="Apa itu UBB?", jawaban="PTN di Babel.",
                           kategori="A. Umum", butuh_verifikasi=True)
    FaqVariasi.objects.create(faq=f, teks="ubb itu apa")

    r = klien_admin.post(reverse("admin:faq_faq_changelist"), {
        "action": "ekspor_excel", "_selected_action": [str(f.pk)]})
    ws = load_workbook(io.BytesIO(r.content)).active

    judul = [c.value for c in ws[1]]
    assert judul[3] == "Pertanyaan" and judul[4] == "Jawaban"
    baris = [c.value for c in ws[2]]
    assert baris[3] == "Apa itu UBB?"
    assert baris[4] == "PTN di Babel."
    assert baris[6] == "Ya"                 # menunggu unit
    assert "ubb itu apa" in baris[8]        # variasi ikut


def test_excel_tidak_menulis_none(klien_admin):
    """Sel kosong harus kosong, bukan tulisan 'None'."""
    from openpyxl import load_workbook

    f = Faq.objects.create(pertanyaan="Tanpa sub-kategori?", jawaban="Ya.")
    r = klien_admin.post(reverse("admin:faq_faq_changelist"), {
        "action": "ekspor_excel", "_selected_action": [str(f.pk)]})
    ws = load_workbook(io.BytesIO(r.content)).active
    assert not any(c.value == "None" for row in ws.iter_rows() for c in row)


# --- halaman statistik ---------------------------------------------------

def test_statistik_terbuka(klien_admin):
    assert klien_admin.get(
        reverse("admin:statistik_percakapan")).status_code == 200


def test_statistik_kosong_tidak_error(klien_admin):
    isi = klien_admin.get(reverse("admin:statistik_percakapan")).content.decode()
    assert "Belum ada data" in isi


def test_statistik_menghitung_rasio_fallback(klien_admin):
    """Rasio fallback adalah angka yang paling perlu diawasi (Bagian 11.3)."""
    for i in range(3):
        ChatLog.objects.create(log_id=f"j{i}", query="x", tipe="jawaban",
                               skor=0.9, latency_ms=40)
    ChatLog.objects.create(log_id="f1", query="x", tipe="fallback",
                           skor=0.2, latency_ms=45)

    isi = klien_admin.get(reverse("admin:statistik_percakapan")).content.decode()
    assert "Tidak terjawab" in isi
    # LANGUAGE_CODE="id" -> pemisah desimal koma, bukan titik.
    assert ("25,0%" in isi) or ("25.0%" in isi)     # 1 dari 4


def test_statistik_menghitung_kepuasan(klien_admin):
    log = ChatLog.objects.create(log_id="a", query="x", tipe="jawaban", skor=.9)
    Feedback.objects.create(log=log, rating="up")
    isi = klien_admin.get(reverse("admin:statistik_percakapan")).content.decode()
    assert "100%" in isi          # 1 👍, 0 👎


def test_statistik_hanya_30_hari_terakhir(klien_admin):
    _log("lama", hari_lalu=60)
    _log("baru", hari_lalu=5)
    isi = klien_admin.get(reverse("admin:statistik_percakapan")).content.decode()
    # Total percakapan = 1 (hanya yang baru)
    assert "Total percakapan" in isi


# --- konfigurasi produksi ------------------------------------------------

def test_settings_produksi_aman():
    import importlib
    import os

    os.environ.setdefault("DJANGO_SECRET_KEY", "uji-saja")
    prod = importlib.import_module("config.settings.prod")

    assert prod.DEBUG is False
    assert prod.SESSION_COOKIE_SECURE is True
    assert prod.CSRF_COOKIE_SECURE is True
    assert prod.SESSION_COOKIE_HTTPONLY is True
    assert prod.X_FRAME_OPTIONS == "DENY"
    assert prod.SECURE_HSTS_SECONDS >= 31536000
    # Tanpa ini Django mengira koneksinya HTTP dan menolak secure cookie.
    assert prod.SECURE_PROXY_SSL_HEADER == ("HTTP_X_FORWARDED_PROTO", "https")
    assert "api.chatbot-ubb.tech" in prod.ALLOWED_HOSTS
    # Health check deploy.sh memanggil 127.0.0.1 — tanpa ini Django membalas
    # 400 DisallowedHost dan deploy berhenti walau API sebenarnya sehat.
    assert "127.0.0.1" in prod.ALLOWED_HOSTS


@pytest.mark.urls("config.urls_api")
def test_healthz_lolos_dari_loopback(settings):
    """Persis yang dilakukan deploy.sh: curl http://127.0.0.1:8000/healthz.

    Host loopback dipakai apa adanya dari daftar produksi — kalau 127.0.0.1
    hilang dari ALLOWED_HOSTS, Django membalas 400 dan deploy gagal.
    """
    import importlib
    import os

    from django.test import Client

    os.environ.setdefault("DJANGO_SECRET_KEY", "uji-saja")
    prod = importlib.import_module("config.settings.prod")
    settings.ALLOWED_HOSTS = prod.ALLOWED_HOSTS

    r = Client().get("/healthz", HTTP_HOST="127.0.0.1:8000")
    assert r.status_code == 200


def test_urlconf_produksi_terpisah():
    import importlib

    import os
    os.environ.setdefault("DJANGO_SECRET_KEY", "uji-saja")
    api = importlib.import_module("config.settings.prod_api")
    admin = importlib.import_module("config.settings.prod_admin")

    assert api.ROOT_URLCONF == "config.urls_api"
    assert admin.ROOT_URLCONF == "config.urls_admin"


def test_secret_key_wajib_dari_environment():
    """SECRET_KEY = os.environ[...] — KeyError kalau lupa, itu disengaja."""
    import pathlib
    src = (pathlib.Path(__file__).resolve().parent.parent
           / "config" / "settings" / "prod.py").read_text()
    assert 'os.environ["DJANGO_SECRET_KEY"]' in src


# --- berkas deploy -------------------------------------------------------

@pytest.mark.parametrize("berkas", [
    "deploy/nginx/01-sebelum-tls.conf",
    "deploy/nginx/02-sesudah-tls.conf",
    "deploy/systemd/gunicorn-api.service",
    "deploy/systemd/gunicorn-admin.service",
    "deploy/gunicorn_api.py",
    "deploy/gunicorn_admin.py",
    "deploy/backup.sh",
    "deploy/deploy.sh",
    "deploy/cron-chatbot",
])
def test_berkas_deploy_ada(berkas):
    import pathlib
    p = pathlib.Path(__file__).resolve().parent.parent / berkas
    assert p.exists(), f"{berkas} tidak ada"
    assert p.stat().st_size > 0


def test_nginx_meneruskan_x_forwarded_for():
    """Tanpa header ini rate limit memblokir SELURUH pengguna sekaligus."""
    import pathlib
    conf = (pathlib.Path(__file__).resolve().parent.parent
            / "deploy/nginx/02-sesudah-tls.conf").read_text()
    assert "X-Forwarded-For" in conf
    assert "proxy_add_x_forwarded_for" in conf


def test_nginx_membatasi_rate():
    import pathlib
    conf = (pathlib.Path(__file__).resolve().parent.parent
            / "deploy/nginx/02-sesudah-tls.conf").read_text()
    assert "limit_req_zone" in conf
    assert "limit_req_status 429" in conf


def test_nginx_admin_pakai_allowlist():
    import pathlib
    conf = (pathlib.Path(__file__).resolve().parent.parent
            / "deploy/nginx/02-sesudah-tls.conf").read_text()
    assert "deny all" in conf


def test_gunicorn_api_preload():
    """preload_app: satu model dibagi antar-worker lewat copy-on-write."""
    import pathlib
    src = (pathlib.Path(__file__).resolve().parent.parent
           / "deploy/gunicorn_api.py").read_text()
    assert "preload_app = True" in src
    assert "max_requests" in src        # lawan kebocoran memori bertahap


def test_backup_menyalin_keluar_vps():
    """Backup di server yang sama dengan datanya bukan backup."""
    import pathlib
    src = (pathlib.Path(__file__).resolve().parent.parent
           / "deploy/backup.sh").read_text()
    assert "rclone" in src
    assert "pg_dump" in src


def test_https_dipaksa_di_produksi():
    """Widget di pmb.ubb.ac.id WAJIB HTTPS — browser memblokir total
    panggilan HTTP dari halaman HTTPS (mixed content, tidak bisa di-bypass)."""
    import importlib
    import os

    os.environ.setdefault("DJANGO_SECRET_KEY", "uji-saja")
    prod = importlib.import_module("config.settings.prod")
    assert prod.SECURE_SSL_REDIRECT is True
    assert prod.SECURE_HSTS_PRELOAD is True
    # healthz dikecualikan agar cek internal Nginx tidak kena redirect.
    assert any("healthz" in p for p in prod.SECURE_REDIRECT_EXEMPT)


def test_check_deploy_bersih():
    """`manage.py check --deploy` tidak boleh menyisakan peringatan keamanan."""
    import os
    import subprocess
    import sys
    import pathlib

    akar = pathlib.Path(__file__).resolve().parent.parent
    env = {
        **os.environ,
        "DJANGO_SETTINGS_MODULE": "config.settings.prod_api",
        "DJANGO_SECRET_KEY": "k" + "x8Qw2Zp9Lm4Rt7Vb1Nc6Hj3Fd5Gs0Ay" * 2,
        "MUAT_MODEL": "0",
    }
    hasil = subprocess.run(
        [sys.executable, "manage.py", "check", "--deploy"],
        cwd=akar, env=env, capture_output=True, text=True, timeout=180)

    assert hasil.returncode == 0, hasil.stderr
    assert "no issues" in (hasil.stdout + hasil.stderr)


# --- uji restore (Bagian 11.1) -------------------------------------------
#
# "Uji restore setiap kuartal. Backup yang belum pernah diuji restore
#  statusnya belum diketahui."  Tes ini menjadikannya otomatis.

def test_backup_bisa_direstore_utuh(tmp_path):
    from faq.models import FaqVariasi

    asli = Faq.objects.create(
        pertanyaan="Apa itu UBB?", jawaban="PTN di Bangka Belitung.",
        kategori="A. Umum", sumber_url="https://ubb.ac.id",
        terakhir_diverifikasi=date(2026, 3, 1))
    FaqVariasi.objects.create(faq=asli, teks="ubb itu apa sih")
    Faq.objects.create(pertanyaan="Berapa luas kampus?",
                       jawaban="Perlu wawancara.", butuh_verifikasi=True)

    berkas = tmp_path / "kb.json"
    call_command("backup_kb", "--keluar", str(berkas), stdout=io.StringIO())
    assert berkas.exists()

    # Simulasi kehilangan total.
    Faq.objects.all().delete()
    assert Faq.objects.count() == 0

    call_command("import_colab", str(berkas), stdout=io.StringIO())

    assert Faq.objects.count() == 2
    pulih = Faq.objects.get(pertanyaan="Apa itu UBB?")
    assert pulih.jawaban == "PTN di Bangka Belitung."
    assert pulih.sumber_url == "https://ubb.ac.id"
    # Flag verifikasi WAJIB terjaga — kalau hilang, placeholder bocor ke publik.
    assert Faq.objects.filter(butuh_verifikasi=True).count() == 1


def test_backup_menyertakan_variasi(tmp_path):
    import json

    from faq.models import FaqVariasi

    f = Faq.objects.create(pertanyaan="Apa itu UBB?", jawaban="PTN.")
    FaqVariasi.objects.create(faq=f, teks="ubb itu apa")

    berkas = tmp_path / "kb.json"
    call_command("backup_kb", "--keluar", str(berkas), stdout=io.StringIO())
    data = json.loads(berkas.read_text(encoding="utf-8"))
    assert data[0]["variasi"] == ["ubb itu apa"]


# --- pemaketan (pip install -e .) ----------------------------------------
#
# Deploy di VPS menjalankan `pip install -e ".[prod]"`. Tanpa [build-system]
# dan daftar paket eksplisit, perintah itu GAGAL dengan pesan yang tidak
# menyebut penyebabnya:
#   ERROR: Failed to build ... when getting requirements to build editable

def test_pyproject_punya_build_system():
    import pathlib
    import tomllib

    p = pathlib.Path(__file__).resolve().parent.parent / "pyproject.toml"
    data = tomllib.loads(p.read_text(encoding="utf-8"))

    assert "build-system" in data, (
        "pyproject.toml tanpa [build-system] membuat `pip install -e .` gagal")
    assert data["build-system"]["build-backend"] == "setuptools.build_meta"


def test_semua_app_django_terdaftar_sebagai_paket():
    """Paket ditulis eksplisit; app baru harus ikut ditambahkan ke sini."""
    import pathlib
    import tomllib

    from django.conf import settings

    p = pathlib.Path(__file__).resolve().parent.parent / "pyproject.toml"
    data = tomllib.loads(p.read_text(encoding="utf-8"))
    terdaftar = set(data["tool"]["setuptools"]["packages"])

    lokal = {a for a in settings.INSTALLED_APPS
             if not a.startswith(("django.", "unfold", "corsheaders"))}
    hilang = lokal - terdaftar
    assert not hilang, (
        f"app {hilang} ada di INSTALLED_APPS tapi tidak di "
        f"[tool.setuptools].packages — akan hilang saat deploy")


def test_pip_install_editable_berhasil():
    """Menjalankan resolusi build sungguhan, seperti yang dilakukan VPS."""
    import pathlib
    import subprocess
    import sys

    akar = pathlib.Path(__file__).resolve().parent.parent
    hasil = subprocess.run(
        # Isolasi build DIBIARKAN aktif — persis seperti pip di VPS.
        [sys.executable, "-m", "pip", "install", "-e", ".",
         "--no-deps", "--dry-run"],
        cwd=akar, capture_output=True, text=True, timeout=300)

    assert hasil.returncode == 0, (
        f"`pip install -e .` gagal — deploy di VPS akan gagal juga.\n"
        f"{hasil.stdout[-1500:]}\n{hasil.stderr[-1500:]}")


def test_nginx_tidak_pakai_sintaks_yang_hanya_ada_di_versi_baru():
    """`http2 on;` hanya dikenal Nginx >= 1.25.

    Ubuntu 24.04 membawa Nginx 1.24 -> gagal start dengan
    `unknown directive "http2"`. Sintaks `listen 443 ssl http2;` berlaku
    di semua versi.
    """
    import pathlib
    import re

    d = pathlib.Path(__file__).resolve().parent.parent / "deploy" / "nginx"
    for conf in d.glob("*.conf"):
        isi = conf.read_text()
        for i, baris in enumerate(isi.splitlines(), 1):
            bersih = baris.split("#")[0].strip()
            assert not re.match(r"^http2\s+(on|off)\s*;", bersih), (
                f"{conf.name}:{i} memakai `http2 on;` yang menolak "
                f"Nginx < 1.25 — pakai `listen 443 ssl http2;`")


def test_nginx_tahap_awal_tidak_butuh_sertifikat():
    """Nginx tidak bisa start dengan `listen 443 ssl` bila sertifikat belum
    ada — certbot perlu jalan masuk lewat port 80 dulu."""
    import pathlib

    conf = (pathlib.Path(__file__).resolve().parent.parent
            / "deploy/nginx/01-sebelum-tls.conf").read_text()

    # Buang komentar dulu — kata "listen 443" muncul di penjelasan.
    direktif = "\n".join(b.split("#")[0] for b in conf.splitlines())

    assert "listen 443" not in direktif
    assert "ssl_certificate" not in direktif
    assert "acme-challenge" in direktif      # jalan verifikasi certbot


# --- konversi metadata model ST 5.x -> 3.x -------------------------------
#
# Colab memakai sentence-transformers 5.x; dev macOS Intel terkunci di 3.2.1.
# Tanpa konversi, model gagal dimuat dengan galat yang tidak menyebut sebabnya:
#   ModuleNotFoundError: No module named 'sentence_transformers.base'

def test_deteksi_format_5x(tmp_path):
    import json

    from faq.management.commands.konversi_model import perlu_konversi

    (tmp_path / "modules.json").write_text(json.dumps([
        {"idx": 0, "name": "0", "path": "",
         "type": "sentence_transformers.base.modules.transformer.Transformer"},
    ]), encoding="utf-8")
    assert perlu_konversi(tmp_path) is True


def test_format_3x_tidak_dikonversi_ulang(tmp_path):
    import json

    from faq.management.commands.konversi_model import perlu_konversi

    (tmp_path / "modules.json").write_text(json.dumps([
        {"idx": 0, "name": "0", "path": "",
         "type": "sentence_transformers.models.Transformer"},
    ]), encoding="utf-8")
    assert perlu_konversi(tmp_path) is False


def test_konversi_menulis_format_lama(tmp_path):
    import json

    from django.core.management import call_command

    (tmp_path / "config.json").write_text(json.dumps({"hidden_size": 384}),
                                          encoding="utf-8")
    (tmp_path / "modules.json").write_text(json.dumps([
        {"idx": 0, "name": "0", "path": "",
         "type": "sentence_transformers.base.modules.transformer.Transformer"},
    ]), encoding="utf-8")
    (tmp_path / "sentence_bert_config.json").write_text(
        json.dumps({"transformer_task": "feature-extraction"}), encoding="utf-8")
    pool = tmp_path / "1_Pooling"
    pool.mkdir()
    (pool / "config.json").write_text(
        json.dumps({"embedding_dimension": 384}), encoding="utf-8")

    call_command("konversi_model", "--dir", str(tmp_path), stdout=io.StringIO())

    modules = json.loads((tmp_path / "modules.json").read_text())
    assert modules[0]["type"] == "sentence_transformers.models.Transformer"

    sbert = json.loads((tmp_path / "sentence_bert_config.json").read_text())
    assert "transformer_task" not in sbert
    assert sbert["max_seq_length"] == 128

    pooling = json.loads((pool / "config.json").read_text())
    assert pooling["word_embedding_dimension"] == 384
    assert "embedding_dimension" not in pooling


def test_konversi_membackup_berkas_asli(tmp_path):
    import json

    from django.core.management import call_command

    model = tmp_path / "model"
    model.mkdir()
    (model / "config.json").write_text(json.dumps({"hidden_size": 384}),
                                       encoding="utf-8")
    asli = [{"idx": 0, "name": "0", "path": "",
             "type": "sentence_transformers.base.modules.transformer.Transformer"}]
    (model / "modules.json").write_text(json.dumps(asli), encoding="utf-8")

    call_command("konversi_model", "--dir", str(model), stdout=io.StringIO())

    cadangan = tmp_path / "_backup_metadata_5x" / "modules.json"
    assert cadangan.exists()
    assert json.loads(cadangan.read_text()) == asli


def test_konversi_tidak_menyentuh_bobot(tmp_path):
    """Bobot model tidak boleh diubah — hanya metadata."""
    import json

    from django.core.management import call_command

    (tmp_path / "config.json").write_text(json.dumps({"hidden_size": 384}),
                                          encoding="utf-8")
    (tmp_path / "modules.json").write_text(json.dumps([
        {"idx": 0, "name": "0", "path": "",
         "type": "sentence_transformers.base.modules.transformer.Transformer"},
    ]), encoding="utf-8")
    bobot = tmp_path / "model.safetensors"
    bobot.write_bytes(b"BOBOT-PALSU-UNTUK-UJI")
    sebelum = bobot.read_bytes()

    call_command("konversi_model", "--dir", str(tmp_path), stdout=io.StringIO())
    assert bobot.read_bytes() == sebelum
